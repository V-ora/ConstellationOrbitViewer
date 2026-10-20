# -*- coding: utf-8 -*-
"""生成《千帆星座轨道状态核查》工作簿
   明细(257颗) + 批次汇总(公式联动) + 说明
"""
import csv, os, datetime

try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation


def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value


XL_HEAD = xl_color("#4472C4")
XL_HEAD_TXT = xl_color("#FFFFFF")
XL_BORDER = xl_color("#BFBFBF")
XL_TITLE = xl_color("#2F5597")
XL_NOTE = xl_color("#F2F6FC")

SRC = r"D:/Work Buddy 工作区/qianfan_orbit/千帆星座轨道状态核查.csv"
OUT = r"D:/Work Buddy 工作区/千帆星座轨道状态核查.xlsx"

# ---------------- 读取源数据 ----------------
rows = list(csv.reader(open(SRC, encoding="utf-8-sig")))
head, data = rows[0], [r for r in rows[1:] if r and r[0].strip()]
assert len(data) == 257, len(data)
NCOL = len(head)

DATE_COLS = {"发射日期", "开始抬升日期", "抵达工作轨道日期", "数据起始", "数据截止"}
INT_COLS = {"NORAD编号", "批次序号", "抬升历时天", "发射至就位天数", "已保持天数",
            "平台保持天数", "历史数据点数", "运行状态码"}
F1_COLS = {"当前平均高度km", "入轨停泊高度km", "近60天平台高度km", "当前远地点km",
           "当前近地点km"}
SIGN_COLS = {"距工作轨道km", "近30天变化km"}


def conv(col, v):
    v = (v or "").strip()
    if v == "":
        return None
    if col in DATE_COLS:
        return datetime.date.fromisoformat(v)
    if col in INT_COLS:
        try:
            return int(float(v))
        except ValueError:
            return v
    if col in F1_COLS or col in SIGN_COLS:
        try:
            return float(v)
        except ValueError:
            return v
    return v


wb = Workbook()
wb.properties.title = "千帆星座轨道状态核查"
wb.calculation.fullCalcOnLoad = True

thin = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FONT = Font(name="微软雅黑", size=10, bold=True, color=XL_HEAD_TXT)
HEAD_FILL = PatternFill("solid", fgColor=XL_HEAD)
BODY_FONT = Font(name="微软雅黑", size=10)
HEAD_AL = Alignment(horizontal="center", vertical="center", wrap_text=True)
CEN = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")

# ================= Sheet1 明细 =================
ws = wb.active
ws.title = "轨道状态核查"
# 锚点：第1行表头，第2~258行数据
cols = ["序号"] + head                                   # 29 列 A..AC
for c, name in enumerate(cols, 1):
    cell = ws.cell(row=1, column=c, value=name)
    cell.font, cell.fill, cell.alignment, cell.border = HEAD_FONT, HEAD_FILL, HEAD_AL, BORDER
for i, r in enumerate(data, 1):
    row = 1 + i
    for c, (name, raw) in enumerate(zip(head, r), 2):
        cell = ws.cell(row=row, column=c, value=conv(name, raw))
        cell.font, cell.border = BODY_FONT, BORDER
        if name in DATE_COLS:
            cell.number_format = "yyyy-mm-dd"; cell.alignment = CEN
        elif name in INT_COLS:
            cell.number_format = "0"; cell.alignment = CEN
        elif name in F1_COLS:
            cell.number_format = "0.0"; cell.alignment = RIGHT
        elif name in SIGN_COLS:
            cell.number_format = "+0.0;-0.0;0.0"; cell.alignment = RIGHT
        elif name in ("轨道周期min", "倾角deg"):
            cell.number_format = "0.00"; cell.alignment = RIGHT
        elif name in ("轨道状态", "所处阶段", "发射批次"):
            cell.alignment = CEN
        else:
            cell.alignment = LEFT
    sc = ws.cell(row=row, column=1, value=i)
    sc.font, sc.border, sc.alignment, sc.number_format = BODY_FONT, BORDER, CEN, "0"
    ws.cell(row=row, column=cols.index("轨道状态") + 1).alignment = CEN
    ws.cell(row=row, column=cols.index("所处阶段") + 1).alignment = CEN

LAST = 1 + len(data)
STATUS_COL = get_column_letter(cols.index("轨道状态") + 1)
BATCH_COL = get_column_letter(cols.index("批次序号") + 1)
ALT_COL = get_column_letter(cols.index("当前平均高度km") + 1)
CLIMB_COL = get_column_letter(cols.index("抬升历时天") + 1)
DELTA_COL = get_column_letter(cols.index("距工作轨道km") + 1)
STATUS_LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
LASTCOL = STATUS_LETTERS[NCOL] if NCOL < 26 else "AC"

# 列宽
WIDTH = {"序号": 6, "卫星名称": 20, "NORAD编号": 11, "国际编号": 12, "发射日期": 12,
         "批次序号": 8, "发射批次": 19, "发射场": 9, "运行状态码": 11,
         "当前远地点km": 13, "当前近地点km": 13, "当前平均高度km": 15,
         "轨道周期min": 12, "倾角deg": 10, "入轨停泊高度km": 15, "开始抬升日期": 13,
         "抵达工作轨道日期": 17, "抬升历时天": 12, "发射至就位天数": 15,
         "已保持天数": 12, "近60天平台高度km": 17, "平台保持天数": 14,
         "距工作轨道km": 14, "近30天变化km": 14, "轨道状态": 22, "所处阶段": 18,
         "历史数据点数": 13, "数据起始": 12, "数据截止": 12}
for c, name in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(c)].width = WIDTH.get(name, 12)
ws.row_dimensions[1].height = 34

ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{LASTCOL}{LAST}"

# 条件格式：轨道状态五色 + 距工作轨道为负提示
STATUS_STYLE = [("已抵达工作轨道", "#C6EFCE", "#006100"),
                ("已进入工作轨道（在轨调整）", "#DAF3F8", "#0B6E7F"),
                ("爬升中", "#FFC7CE", "#9C0006"),
                ("入轨停泊（待抬升）", "#FFE8CC", "#9C6500"),
                ("抬升中止（未达工作轨道）", "#E4DAF7", "#5B3E9C")]
rng = f"{STATUS_COL}2:{STATUS_COL}{LAST}"
for txt, bg, fg in STATUS_STYLE:
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'EXACT(${STATUS_COL}2,"{txt}")'],
        fill=PatternFill(patternType="solid", fgColor=xl_color(bg), bgColor=xl_color(bg)),
        font=Font(name="微软雅黑", size=10, color=xl_color(fg), bold=True), stopIfTrue=False))
ws.conditional_formatting.add(f"{DELTA_COL}2:{DELTA_COL}{LAST}", CellIsRule(
    operator="lessThan", formula=["-3"],
    font=Font(name="微软雅黑", size=10, color=xl_color("#9C0006"), bold=True)))

dv = DataValidation(type="list", formula1='"%s"' % ",".join(s[0] for s in STATUS_STYLE),
                    allow_blank=True, showDropDown=False)
ws.add_data_validation(dv)
dv.add(f"{STATUS_COL}2:{STATUS_COL}{LAST}")

# ================= Sheet2 批次汇总 =================
s2 = wb.create_sheet("批次汇总")
batches = sorted({(int(r[head.index("批次序号")]), r[head.index("发射批次")],
                   r[head.index("发射日期")]) for r in data}, key=lambda x: x[0])
BH = ["批次", "发射日期", "发射数量", "已抵达工作轨道", "在轨调整中", "爬升中",
      "入轨停泊待抬升", "抬升中止未达标", "已就位合计", "就位率", "平均当前高度km",
      "平均抬升历时天"]
s2["A1"] = "千帆星座发射批次轨道就位情况汇总"
s2["A1"].font = Font(name="微软雅黑", size=13, bold=True, color=XL_TITLE)
s2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(BH))
s2["A1"].alignment = Alignment(horizontal="left", vertical="center")
s2.row_dimensions[1].height = 26
for c, name in enumerate(BH, 1):
    cell = s2.cell(row=2, column=c, value=name)
    cell.font, cell.fill, cell.alignment, cell.border = HEAD_FONT, HEAD_FILL, HEAD_AL, BORDER
s2.row_dimensions[2].height = 32

S = "'轨道状态核查'"
for i, (bid, bname, bdate) in enumerate(batches):
    r = 3 + i
    s2.cell(row=r, column=1, value=bname).alignment = CEN
    s2.cell(row=r, column=2, value=datetime.date.fromisoformat(bdate)).number_format = "yyyy-mm-dd"
    s2.cell(row=r, column=2).alignment = CEN
    s2.cell(row=r, column=3, value=f"=COUNTIF({S}!${BATCH_COL}:${BATCH_COL},$M{r})")
    s2.cell(row=r, column=4, value=f'=COUNTIFS({S}!${BATCH_COL}:${BATCH_COL},$M{r},{S}!${STATUS_COL}:${STATUS_COL},"已抵达工作轨道")')
    s2.cell(row=r, column=5, value=f'=COUNTIFS({S}!${BATCH_COL}:${BATCH_COL},$M{r},{S}!${STATUS_COL}:${STATUS_COL},"已进入工作轨道（在轨调整）")')
    s2.cell(row=r, column=6, value=f'=COUNTIFS({S}!${BATCH_COL}:${BATCH_COL},$M{r},{S}!${STATUS_COL}:${STATUS_COL},"爬升中")')
    s2.cell(row=r, column=7, value=f'=COUNTIFS({S}!${BATCH_COL}:${BATCH_COL},$M{r},{S}!${STATUS_COL}:${STATUS_COL},"入轨停泊（待抬升）")')
    s2.cell(row=r, column=8, value=f'=COUNTIFS({S}!${BATCH_COL}:${BATCH_COL},$M{r},{S}!${STATUS_COL}:${STATUS_COL},"抬升中止（未达工作轨道）")')
    s2.cell(row=r, column=9, value=f"=D{r}+E{r}")
    s2.cell(row=r, column=10, value=f'=IF(OR(C{r}="",C{r}=0),"",(D{r}+E{r})/C{r})')
    s2.cell(row=r, column=11, value=f'=IFERROR(ROUND(AVERAGEIF({S}!${BATCH_COL}:${BATCH_COL},$M{r},{S}!${ALT_COL}:${ALT_COL}),1),"")')
    s2.cell(row=r, column=12, value=f'=IFERROR(ROUND(AVERAGEIFS({S}!${CLIMB_COL}:${CLIMB_COL},{S}!${BATCH_COL}:${BATCH_COL},$M{r},{S}!${CLIMB_COL}:${CLIMB_COL},">0"),0),"")')
    for c in range(1, len(BH) + 1):
        cell = s2.cell(row=r, column=c)
        cell.border = BORDER
        if cell.font.name != "微软雅黑" or cell.font.size != 10:
            cell.font = BODY_FONT
    for c in (3, 4, 5, 6, 7, 8, 9):
        s2.cell(row=r, column=c).number_format = "0"
        s2.cell(row=r, column=c).alignment = CEN
    s2.cell(row=r, column=10).number_format = "0.0%"
    s2.cell(row=r, column=10).alignment = CEN
    s2.cell(row=r, column=11).number_format = "0.0"
    s2.cell(row=r, column=12).number_format = "0"
    s2.cell(row=r, column=1).font = Font(name="微软雅黑", size=10, bold=True)
    # 序号列写 批次序号 供 COUNTIF 匹配
    s2.cell(row=r, column=13, value=bid).font = Font(name="微软雅黑", size=8, color=xl_color("#FFFFFF"))

TOT = 3 + len(batches)
s2.cell(row=TOT, column=1, value="合计").font = Font(name="微软雅黑", size=10, bold=True)
for c in range(2, 10):
    L = get_column_letter(c)
    cell = s2.cell(row=TOT, column=c, value=f"=SUM({L}3:{L}{TOT - 1})" if c >= 3 else None)
    cell.font = Font(name="微软雅黑", size=10, bold=True)
    cell.number_format = "0"
    cell.alignment = CEN
s2.cell(row=TOT, column=10, value=f'=IF(OR(C{TOT}="",C{TOT}=0),"",(D{TOT}+E{TOT})/C{TOT})')
s2.cell(row=TOT, column=10).number_format = "0.0%"
s2.cell(row=TOT, column=10).alignment = CEN
s2.cell(row=TOT, column=11, value=f'=IFERROR(ROUND(AVERAGE({S}!${ALT_COL}:${ALT_COL}),1),"")')
s2.cell(row=TOT, column=11).number_format = "0.0"
s2.cell(row=TOT, column=12, value=f'=IFERROR(ROUND(AVERAGEIF({S}!${CLIMB_COL}:${CLIMB_COL},">0"),0),"")')
s2.cell(row=TOT, column=12).number_format = "0"
for c in range(1, 13):
    cell = s2.cell(row=TOT, column=c)
    cell.border = Border(left=thin, right=thin, top=Side(style="medium", color=XL_BORDER), bottom=thin)
    cell.fill = PatternFill("solid", fgColor=xl_color("#D9E2F3"))
    if c >= 2:
        cell.font = Font(name="微软雅黑", size=10, bold=True)

for c, w in zip(range(1, 13), [19, 12, 10, 15, 11, 9, 15, 15, 11, 10, 15, 15]):
    s2.column_dimensions[get_column_letter(c)].width = w
s2.freeze_panes = "A3"

note = s2.cell(row=TOT + 2, column=1,
               value="说明：本表所有指标均由“轨道状态核查”明细表用公式实时汇总（COUNTIF / COUNTIFS / AVERAGEIF），"
                     "修改明细表后自动更新；第13列隐藏列为批次序号，供公式匹配使用。")
note.font = Font(name="微软雅黑", size=9, color=xl_color("#7F7F7F"))
note.alignment = Alignment(horizontal="left", vertical="center")
s2.cell(row=2, column=13, value="批次序号（辅助列）").font = Font(name="微软雅黑", size=8, color=xl_color("#BFBFBF"))
s2.column_dimensions["M"].hidden = True

# ================= Sheet3 说明 =================
s3 = wb.create_sheet("说明")
s3["A1"] = "千帆星座（Qianfan / G60）轨道状态核查 —— 数据说明"
s3["A1"].font = Font(name="微软雅黑", size=14, bold=True, color=XL_TITLE)
s3.merge_cells("A1:C1")
s3.row_dimensions[1].height = 28
ITEMS = [
    ("数据来源", "CelesTrak（https://celestrak.org）SATCAT 目录记录 + 逐星历史轨道演化序列（GP 历史，graph-orbit-data）"),
    ("目标星座", "千帆星座 / Qianfan（G60 星链），CelesTrak 名称前缀 QIANFAN"),
    ("卫星数量", f"{len(data)} 颗在轨载荷，分属 {len(batches)} 个发射批次（2024-08-06 至 2026-09-15）"),
    ("数据截止", "2026-09-27（轨道高度历史序列末次数据点 2026-09-26）"),
    ("工作轨道基准", "1068.5 km（近圆轨道，平均高度取远地点与近地点均值；轨道周期 106.59 min，倾角约 88.99°）"),
    ("轨道高度定义", "取 GP 历史序列中的“半长轴高度（Semi-Major Axis Altitude）”，按日取中位数后做 3 日滑动中位数平滑"),
    ("爬升过程", "卫星入轨后先在约 783–1 037 km 高度停泊（多数约 803 km），随后以电推进持续抬升，历时中位约 48 天抵达 1068 km 工作轨道并长期保持"),
    ("状态判定口径", "① 已抵达工作轨道：当前平均高度 ≥ 工作轨道 −3 km，且近 30 天高度变化 < 1.5 km；"
                 "② 已进入工作轨道（在轨调整）：已进入目标高度带但仍在调整；"
                 "③ 爬升中：低于目标带且近期仍持续抬升；"
                 "④ 入轨停泊（待抬升）：新入轨、尚未开始抬升；"
                 "⑤ 抬升中止（未达工作轨道）：长期停留在中间高度且不再抬升"),
    ("关键列说明", "距工作轨道km = 当前平均高度 − 1068.5（负值表示尚未达标）；"
               "抬升历时天 = 抵达工作轨道日期 − 开始抬升日期；"
               "发射至就位天数 = 抵达工作轨道日期 − 发射日期；"
               "平台保持天数 = 当前高度保持不变（波动 < 2.5 km）的持续天数"),
    ("使用提示", "明细表已开启筛选与冻结窗格，可按“轨道状态”“发射批次”筛选；"
             "“批次汇总”表所有数字均由公式联动词明细表，勿手工覆盖"),
    ("数据说明", "轨道高度为公开 TLE/GP 数据的推算结果，与卫星实际在轨状态可能存在偏差；"
             "“抬升中止”为基于高度长期停滞的推断，可能对应电推变轨中止、任务调整或其它工况"),
]
for i, (k, v) in enumerate(ITEMS, 3):
    s3.cell(row=i, column=1, value=k).font = Font(name="微软雅黑", size=10, bold=True,
                                                  color=xl_color("#1F3864"))
    s3.cell(row=i, column=1).alignment = Alignment(horizontal="left", vertical="top")
    s3.cell(row=i, column=1).fill = PatternFill("solid", fgColor=XL_NOTE)
    s3.cell(row=i, column=1).border = BORDER
    c = s3.cell(row=i, column=2, value=v)
    c.font = Font(name="微软雅黑", size=10)
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c.border = BORDER
    s3.row_dimensions[i].height = 46
s3.column_dimensions["A"].width = 18
s3.column_dimensions["B"].width = 108

wb.save(OUT)

# ---------------- 注入公式缓存值 ----------------
# 本机无 LibreOffice / formulas 引擎,为使预览器(pandas/data_only)也能读到数值,
# 这里按明细数据算出汇总结果,写入 xlsx 的 <v> 缓存节点;同时置 fullCalcOnLoad,
# 用户用 Excel/WPS 打开时会自动重算,公式仍可编辑。
import zipfile, re, shutil, tempfile

vals = {}
for i, (bid, bname, bdate) in enumerate(batches):
    r = 3 + i
    rows_b = [x for x in data if int(x[head.index("批次序号")]) == bid]
    def cs(name):
        return sum(1 for x in rows_b if x[head.index("轨道状态")] == name)
    n = len(rows_b)
    d_ok, d_adj = cs("已抵达工作轨道"), cs("已进入工作轨道（在轨调整）")
    alts = [float(x[head.index("当前平均高度km")]) for x in rows_b
            if x[head.index("当前平均高度km")].strip()]
    climbs = [int(x[head.index("抬升历时天")]) for x in rows_b
              if x[head.index("抬升历时天")].strip() and int(x[head.index("抬升历时天")]) > 0]
    vals[f"C{r}"] = n
    vals[f"D{r}"] = d_ok
    vals[f"E{r}"] = d_adj
    vals[f"F{r}"] = cs("爬升中")
    vals[f"G{r}"] = cs("入轨停泊（待抬升）")
    vals[f"H{r}"] = cs("抬升中止（未达工作轨道）")
    vals[f"I{r}"] = d_ok + d_adj
    vals[f"J{r}"] = round((d_ok + d_adj) / n, 4) if n else ""
    vals[f"K{r}"] = round(sum(alts) / len(alts), 1) if alts else ""
    vals[f"L{r}"] = round(sum(climbs) / len(climbs)) if climbs else ""
tot_rows = [x for x in data]
tot_alts = [float(x[head.index("当前平均高度km")]) for x in tot_rows]
tot_climbs = [int(x[head.index("抬升历时天")]) for x in tot_rows
              if x[head.index("抬升历时天")].strip() and int(x[head.index("抬升历时天")]) > 0]
tot_ok = sum(1 for x in tot_rows if x[head.index("轨道状态")].startswith("已抵达")
             or x[head.index("轨道状态")].startswith("已进入"))
for c, key in zip(range(3, 10), "CDEFGHI"):
    vals[f"{key}{TOT}"] = sum(vals[f"{key}{r}"] for r in range(3, TOT))
vals[f"J{TOT}"] = round(tot_ok / len(tot_rows), 4)
vals[f"K{TOT}"] = round(sum(tot_alts) / len(tot_alts), 1)
vals[f"L{TOT}"] = round(sum(tot_climbs) / len(tot_climbs))


def inject_cache(path, sheet_name, values):
    tmp = path + ".tmp"
    zin = zipfile.ZipFile(path)
    names = zin.namelist()
    # 通过 workbook.xml 的 sheet 顺序定位 sheet xml
    wbxml = zin.read("xl/workbook.xml").decode("utf-8")
    order = re.findall(r'<sheet[^>]*name="([^"]+)"', wbxml)
    idx = order.index(sheet_name) + 1
    target = f"xl/worksheets/sheet{idx}.xml"
    xml = zin.read(target).decode("utf-8")
    cnt = 0
    for ref, v in values.items():
        pat = re.compile(r'(<c r="%s"(?P<attr>[^>]*)>)(?P<f><f>[^<]*</f>)' % ref)
        m = pat.search(xml)
        if not m:
            continue
        tag, attr, f = m.group(1), m.group("attr"), m.group("f")
        if isinstance(v, str):
            tag2 = f'<c r="{ref}"{attr} t="str">'
            xml = xml[:m.start()] + tag2 + f + f"<v>{v}</v>" + xml[m.end():]
        else:
            xml = xml[:m.start()] + tag + f + f"<v>{v}</v>" + xml[m.end():]
        cnt += 1
    zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
    for nm in names:
        zout.writestr(nm, xml.encode("utf-8") if nm == target else zin.read(nm))
    zout.close(); zin.close()
    shutil.move(tmp, path)
    return cnt


n_inj = inject_cache(OUT, "批次汇总", vals)
print("saved", OUT, "| sheets:", wb.sheetnames, "| 明细行数:", len(data),
      "| 缓存值注入:", n_inj, "/", len(vals))
